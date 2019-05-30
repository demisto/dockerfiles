from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from sane_doc_reports.domain.CellObject import CellObject
from sane_doc_reports.transform.markdown.MarkdownSection import MarkdownSection
from sane_doc_reports.domain.Wrapper import Wrapper
from sane_doc_reports.conf import DEBUG, MD_TYPE_QUOTE
from sane_doc_reports.elements import markdown, error
from sane_doc_reports.styles.colors import name_to_hex


class QuoteWrapper(Wrapper):

    def wrap(self):
        if DEBUG:
            print('Wrapping quote...')

        self.cell_object.add_paragraph()
        new_cell = self.cell_object.cell.add_table(1, 1).cell(0, 0)

        quote_color = name_to_hex("cornsilk")
        shading_elm_1 = parse_xml(
            f'<w:shd {nsdecls("w")} w:fill="{quote_color}"/>')
        new_cell._tc.get_or_add_tcPr().append(shading_elm_1)

        self.cell_object = CellObject(new_cell)

        contents = self.section.contents
        if isinstance(contents, str):
            temp_section = MarkdownSection('markdown',
                                           [MarkdownSection('span', contents,
                                                            {}, {})]
                                           , {}, {})
        else:
            temp_section = MarkdownSection('markdown',
                                           contents, {}, {})

        markdown.invoke(self.cell_object, temp_section,
                        invoked_from_wrapper=True)


def invoke(cell_object, section):
    if section.type != MD_TYPE_QUOTE:
        section.contents = 'Called blockquote but not blockquote (quote) -' + \
                           f' [{section}]'
        return error.invoke(cell_object, section)

    QuoteWrapper(cell_object, section).wrap()
