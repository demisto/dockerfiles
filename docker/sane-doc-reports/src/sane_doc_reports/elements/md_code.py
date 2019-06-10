from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from sane_doc_reports.domain.CellObject import CellObject
from sane_doc_reports.styles.utils import insert_cell_background
from sane_doc_reports.transform.markdown.MarkdownSection import MarkdownSection
from sane_doc_reports.conf import DEBUG, MD_TYPE_CODE
from sane_doc_reports.elements import markdown, error
from sane_doc_reports.styles.colors import name_to_hex
from sane_doc_reports.domain.Wrapper import Wrapper


class CodeWrapper(Wrapper):

    def wrap(self):
        if DEBUG:
            print("Wrapping code...")

        if 'inline' not in self.section.extra:
            self.cell_object.add_paragraph()
            # TODO: remove newlines from OXML

        cell = self.cell_object.cell.add_table(1, 1).cell(0, 0)
        code_color = name_to_hex("whitesmoke")
        new_cell = insert_cell_background(cell, code_color)
        self.cell_object = CellObject(new_cell)

        contents = self.section.contents
        if isinstance(contents, str):
            temp_section = MarkdownSection('markdown',
                                           [MarkdownSection('span', contents,
                                                            {}, {})]
                                           , {}, {})
        else:
            temp_section = MarkdownSection('markdown',
                                           contents, {}, {})

        markdown.invoke(self.cell_object, temp_section,
                        invoked_from_wrapper=True)


def invoke(cell_object, section):
    if section.type != MD_TYPE_CODE:
        section.contents = f'Called code but not code -  [{section}]'
        return error.invoke(cell_object, section)

    CodeWrapper(cell_object, section).wrap()
